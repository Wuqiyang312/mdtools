"""文档转换核心逻辑封装"""

import os
import re
from pathlib import Path
from typing import Optional

import pdfplumber
import markdown
from weasyprint import HTML, CSS
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def pdf_to_text(pdf_path: str) -> str:
    """从 PDF 提取文本内容"""
    if not os.path.exists(pdf_path):
        raise FileNotFoundError(f"文件不存在：{pdf_path}")
    
    text_content = []
    with pdfplumber.open(pdf_path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            page_text = page.extract_text()
            if page_text:
                text_content.append(f"## 第 {i} 页\n")
                text_content.append(page_text)
                text_content.append("\n")
            
            tables = page.extract_tables()
            for j, table in enumerate(tables, 1):
                if table:
                    text_content.append(f"**表格 {j}**\n")
                    for row in table:
                        cleaned_row = [
                            str(cell).strip() if cell is not None else ""
                            for cell in row
                        ]
                        text_content.append("| " + " | ".join(cleaned_row) + " |")
                    text_content.append("\n")
    
    return "\n".join(text_content)


def convert_pdf_to_md(pdf_path: str, output_path: Optional[str] = None) -> str:
    """将 PDF 转换为 Markdown"""
    if not os.path.exists(pdf_path):
        raise FileNotFoundError(f"文件不存在：{pdf_path}")
    
    if output_path is None:
        base_name = os.path.splitext(pdf_path)[0]
        output_path = f"{base_name}.md"
    
    content = pdf_to_text(pdf_path)
    md_content = f"# {os.path.basename(pdf_path)}\n\n" + content
    
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(md_content)
    
    return output_path


def get_chinese_css() -> str:
    """获取中文 PDF 样式"""
    return """
    @page {
        size: A4;
        margin: 2.5cm 2cm;
    }
    
    body {
        font-family: "Noto Serif CJK SC", "Source Han Serif CN", "SimSun", serif;
        font-size: 12pt;
        line-height: 1.8;
        text-align: justify;
        color: #333;
    }
    
    h1, h2, h3, h4, h5, h6 {
        font-family: "Noto Sans CJK SC", "Source Han Sans CN", "SimHei", serif;
        margin-top: 1.5em;
        margin-bottom: 0.8em;
        page-break-after: avoid;
    }
    
    h1 { font-size: 18pt; border-bottom: 2px solid #333; padding-bottom: 0.3em; }
    h2 { font-size: 16pt; border-bottom: 1px solid #ddd; padding-bottom: 0.3em; }
    h3 { font-size: 14pt; }
    
    code {
        font-family: "Courier New", Consolas, monospace;
        background-color: #f5f5f5;
        padding: 0.2em 0.4em;
        border-radius: 3px;
    }
    
    pre {
        background-color: #f5f5f5;
        padding: 1em;
        border-radius: 5px;
        overflow-x: auto;
    }
    
    table {
        border-collapse: collapse;
        width: 100%;
        margin: 1em 0;
        font-size: 10pt;
    }
    
    th, td {
        border: 1px solid #ddd;
        padding: 0.5em;
        text-align: left;
    }
    
    th {
        background-color: #f5f5f5;
        font-weight: bold;
    }
    """


def md_to_html(md_path: str) -> str:
    """将 Markdown 转换为 HTML"""
    if not os.path.exists(md_path):
        raise FileNotFoundError(f"文件不存在：{md_path}")
    
    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()
    
    html_content = markdown.markdown(
        md_content, 
        extensions=["tables", "toc", "fenced_code", "codehilite"]
    )
    
    full_html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <title>{os.path.basename(md_path)}</title>
</head>
<body>
{html_content}
</body>
</html>"""
    
    return full_html


def convert_md_to_pdf(
    md_path: str, 
    output_path: Optional[str] = None,
    css_path: Optional[str] = None
) -> str:
    """将 Markdown 转换为 PDF"""
    if not os.path.exists(md_path):
        raise FileNotFoundError(f"文件不存在：{md_path}")
    
    if output_path is None:
        base_name = os.path.splitext(md_path)[0]
        output_path = f"{base_name}.pdf"
    
    html_content = md_to_html(md_path)
    
    if css_path and os.path.exists(css_path):
        css = CSS(filename=css_path)
    else:
        css = CSS(string=get_chinese_css())
    
    html_doc = HTML(string=html_content)
    html_doc.write_pdf(output_path, stylesheets=[css])
    
    return output_path


def convert_md_to_word(md_path: str, output_path: Optional[str] = None) -> str:
    """将 Markdown 转换为 Word"""
    if not os.path.exists(md_path):
        raise FileNotFoundError(f"文件不存在：{md_path}")
    
    if output_path is None:
        base_name = os.path.splitext(md_path)[0]
        output_path = f"{base_name}.docx"
    
    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()
    
    doc = Document()
    
    def set_font_style(paragraph, font_name="微软雅黑", font_size=12):
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            rPr = run._element.get_or_add_rPr()
            rFonts = OxmlElement("w:rFonts")
            rFonts.set(qn("w:eastAsia"), font_name)
            rPr.insert(0, rFonts)
    
    def add_normal_paragraph(text):
        p = doc.add_paragraph(text)
        set_font_style(p, "微软雅黑", 12)
        return p
    
    def add_heading_paragraph(text, level, size_map=None):
        if size_map is None:
            size_map = {1: 24, 2: 18, 3: 16, 4: 14, 5: 12, 6: 12}
        p = doc.add_heading(text, level=level)
        set_font_style(p, "微软雅黑", size_map.get(level, 12))
        return p
    
    lines = md_content.split("\n")
    current_paragraph = []
    
    def flush_paragraph():
        if current_paragraph:
            text = " ".join(current_paragraph)
            if text.strip():
                p = add_normal_paragraph(text)
                p.paragraph_format.space_after = Pt(6)
        current_paragraph.clear()
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        
        if re.match(r"^#{1,6}\s", line):
            flush_paragraph()
            match = re.match(r"^(#{1,6})\s+(.*)", line)
            if match:
                level = len(match.group(1))
                text = match.group(2)
                add_heading_paragraph(text, level)
        
        elif line.startswith("- ") or line.startswith("* "):
            flush_paragraph()
            text = line[2:].strip()
            p = add_normal_paragraph(text)
            p.style = "List Bullet"
        
        elif re.match(r"^\d+\.\s", line):
            flush_paragraph()
            text = re.sub(r"^\d+\.\s", "", line).strip()
            p = add_normal_paragraph(text)
            p.style = "List Number"
        
        elif line.startswith("> "):
            flush_paragraph()
            text = line[2:].strip()
            p = add_normal_paragraph(text)
            p.paragraph_format.left_indent = Cm(1)
            p.italic = True
        
        elif line.startswith("```"):
            flush_paragraph()
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if code_lines:
                code_text = "\n".join(code_lines)
                p = add_normal_paragraph(code_text)
                p.style.font.name = "Courier New"
                p.paragraph_format.left_indent = Cm(0.5)
        
        elif line.strip() == "":
            flush_paragraph()
        
        else:
            current_paragraph.append(line)
        
        i += 1
    
    flush_paragraph()
    doc.save(output_path)
    return output_path
