#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown 转 Word 转换器 (支持中文)
"""

import sys
import os
import re
import argparse
from pathlib import Path
import markdown
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_font_style(paragraph, font_name="微软雅黑", font_size=12):
    """设置段落字体样式"""
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = Pt(font_size)
        rPr = run._element.get_or_add_rPr()
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:eastAsia"), font_name)
        rPr.insert(0, rFonts)


def md_to_word(md_path, output_path=None):
    """将 Markdown 文件转换为 Word 文档"""
    md_path = Path(md_path)
    if not md_path.exists():
        raise FileNotFoundError(f"文件不存在：{md_path}")

    if output_path is None:
        output_path = md_path.with_suffix(".docx")
    else:
        output_path = Path(output_path)

    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    doc = Document()

    def add_normal_paragraph(text):
        p = doc.add_paragraph(text)
        set_font_style(p, "微软雅黑", 12)
        return p

    def add_heading_paragraph(
        text, level, size_map={1: 24, 2: 18, 3: 16, 4: 14, 5: 12, 6: 12}
    ):
        p = doc.add_heading(text, level=level)
        set_font_style(p, "微软雅黑", size_map.get(level, 12))
        return p

    lines = md_content.split("\n")
    current_paragraph = []

    def flush_paragraph():
        if current_paragraph:
            text = " ".join(current_paragraph)
            if text.strip():
                p = add_normal_paragraph(text)
                p.paragraph_format.space_after = Pt(6)
            current_paragraph.clear()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if re.match(r"^#{1,6}\s", line):
            flush_paragraph()
            match = re.match(r"^(#{1,6})\s+(.*)", line)
            if match:
                level = len(match.group(1))
                text = match.group(2)
                add_heading_paragraph(text, level)

        elif line.startswith("- ") or line.startswith("* "):
            flush_paragraph()
            text = line[2:].strip()
            p = add_normal_paragraph(text)
            p.style = "List Bullet"

        elif re.match(r"^\d+\.\s", line):
            flush_paragraph()
            text = re.sub(r"^\d+\.\s", "", line).strip()
            p = add_normal_paragraph(text)
            p.style = "List Number"

        elif line.startswith("> "):
            flush_paragraph()
            text = line[2:].strip()
            p = add_normal_paragraph(text)
            p.paragraph_format.left_indent = Cm(1)
            p.italic = True

        elif line.startswith("```"):
            flush_paragraph()
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if code_lines:
                code_text = "\n".join(code_lines)
                p = add_normal_paragraph(code_text)
                p.style.font.name = "Courier New"
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)

        elif line.strip() == "":
            flush_paragraph()

        else:
            current_paragraph.append(line)

        i += 1

    flush_paragraph()

    doc.save(output_path)
    return output_path


def batch_convert(directory=".", pattern="*.md"):
    """批量转换 Markdown 文件"""
    dir_path = Path(directory)
    md_files = list(dir_path.glob(pattern))

    if not md_files:
        print(f"未找到 Markdown 文件：{dir_path / pattern}")
        return

    print(f"找到 {len(md_files)} 个 Markdown 文件")
    for md_file in md_files:
        if "dist-info" not in str(md_file):
            md_to_word(md_file)


def main():
    parser = argparse.ArgumentParser(
        description="Markdown 转 Word 工具（支持中文）",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  python md2word.py document.md
  python md2word.py document.md -o output.docx
  python md2word.py --batch ./docs/
        """,
    )

    parser.add_argument("md_file", nargs="?", help="Markdown 文件路径")
    parser.add_argument("-o", "--output", help="输出 Word 文件路径")
    parser.add_argument("--batch", nargs="?", const=".", help="批量转换目录")

    args = parser.parse_args()

    if args.batch:
        batch_convert(args.batch)
    elif args.md_file:
        try:
            output_path = md_to_word(args.md_file, args.output)
            print(f"[OK] Conversion successful: {output_path}")
        except Exception as e:
            print(f"[ERROR] Conversion failed: {e}", file=sys.stderr)
            sys.exit(1)
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
